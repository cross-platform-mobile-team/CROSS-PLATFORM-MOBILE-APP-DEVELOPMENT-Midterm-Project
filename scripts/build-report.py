"""Build the TDTU report from one reviewed Markdown source and retained template.

Run with the bundled document Python runtime. Does not modify the reference,
application source, golden baselines or test evidence. Word field refresh and
PDF export are separate steps in render-report.ps1.
"""
from pathlib import Path
import hashlib
import json
import re
import shutil
import zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/report'
OUT = ROOT / 'output/pdf'
LATEX = SOURCE / 'latex'
ASSETS = LATEX / 'assets'
REF = ROOT / 'build/report-template/reference.docx'
REF_HASH = 'b2e75bbbab293cec424726b20dc6880d4a83bc22b5f8a10f4d2b4fd68edefb6d'
TITLE = 'UI Automation Testing in Flutter'
SUBTITLE = 'Implementing End-to-End Testing'
ABSTRACT = (
    'This report investigates UI automation testing through TaskFlow, a Flutter task-management '
    'application with isolated offline, sample and account workspaces. The study connects unit, '
    'widget, golden, native integration and browser tests to specific risks rather than treating '
    'test count as a quality measure. Injectable repositories, clocks and identifiers support '
    'deterministic fixtures, while explicit completion gates expose loading and recovery states. '
    'The account implementation adds a Node.js and SQLite boundary with ownership checks and '
    'revision-based snapshot writes. A deliberate substring-to-prefix search defect is detected '
    'by a permanent regression, with authentic baseline, failure and corrected outputs. Three '
    'rounds of selected test-level commands illustrate local wall-time differences, but unequal '
    'workloads and startup overhead prevent a universal speed ranking. During report preparation, '
    'the updated Flutter suite passed 100 cases and the backend suite passed 14 cases. Retained '
    'Windows and Android native suites each passed 15 cases on 14 and 15 September 2026, '
    'respectively. Android release installation and process-restart persistence were verified '
    'on the local API 36 emulator. Edge and historical Android CI results retain their '
    'revision and platform boundaries. These selected checks do not establish absence of '
    'defects or unconditional release readiness. The study concludes that clear '
    'oracles, isolated state, complementary execution boundaries and honest interpretation are '
    'central to a defensible Flutter testing case study.'
)


def field(p, instruction, display=''):
    run = p.add_run()
    start = OxmlElement('w:fldChar'); start.set(qn('w:fldCharType'), 'begin')
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = ' ' + instruction + ' '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = display
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for item in (start, code, sep, text, end): run._r.append(item)


def diagram(name, boxes, arrows, size=(1400, 760)):
    im = Image.new('RGB', size, 'white'); d = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 28)
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 30)
    for rect, title, lines in boxes:
        d.rounded_rectangle(rect, radius=14, fill='#F8FAFC', outline='#64748B', width=3)
        x1, y1, x2, y2 = rect
        step = min(45, (y2-y1-40)/len([title]+lines))
        yy = y1 + 12
        for i, text in enumerate([title] + lines):
            f = bold if i == 0 else font
            bounds = d.textbbox((0, 0), text, font=f)
            d.text(((x1+x2-bounds[2])/2, yy), text, font=f, fill='#111827')
            yy += step
    for x1, y1, x2, y2 in arrows:
        d.line((x1,y1,x2,y2), fill='#475569', width=4)
        if y2 > y1: pts=[(x2,y2),(x2-9,y2-15),(x2+9,y2-15)]
        elif x2 > x1: pts=[(x2,y2),(x2-15,y2-9),(x2-15,y2+9)]
        else: pts=[(x2,y2),(x2+15,y2-9),(x2+15,y2+9)]
        d.polygon(pts, fill='#475569')
    im.save(ASSETS/name)


def assets():
    ASSETS.mkdir(parents=True, exist_ok=True); OUT.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(REF) as z:
        (ASSETS/'tdtu-logo.png').write_bytes(z.read('word/media/image1.png'))
        inventory = {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
    (ROOT/'build/report-template/package-inventory.json').write_text(json.dumps(inventory, indent=2))
    shutil.copy2(ROOT/'test/golden/baselines/wide-populated.png', ASSETS/'ui-wide.png')
    diagram('testing-layers.png', [
        ((270,20,1130,150),'Native integration and browser',['Real runtime boundaries; selected workflows']),
        ((170,200,1230,330),'Widget and golden',['Interaction states and reviewed visual references']),
        ((60,380,1340,510),'Unit and repository contract',['Validation, transitions, deterministic data behavior'])
    ], [(700,150,700,195),(700,330,700,375)], (1400,540))
    diagram('architecture.png', [
        ((360,10,1040,135),'Presentation',['Task screen, forms and navigation']),
        ((360,190,1040,315),'Application',['TaskController, clock and ID seams']),
        ((360,370,1040,480),'TaskRepository',['load and save contract']),
        ((15,545,445,720),'Local preferences',['Offline records','Real persistence tests']),
        ((485,545,915,720),'In-memory fake',['Isolated sample and QA','Delays and failures']),
        ((955,545,1385,720),'HTTP repository',['Account workspace','Node.js and SQLite'])
    ], [(700,135,700,185),(700,315,700,365),(700,480,230,540),(700,480,700,540),(700,480,1170,540)])
    diagram('data-model.png', [
        ((10,230,425,475),'users',['id primary key','email unique','password and recovery hash','revision']),
        ((560,15,1370,150),'tasks',['(user_id, id) composite key','JSON task payload']),
        ((560,210,1030,390),'sessions',['user_id foreign key','access and refresh digests','expiry values']),
        ((1060,230,1385,365),'used_refresh',['token hash','session_id']),
        ((560,480,1370,630),'audit',['user_id foreign key','event and timestamp; bounded history'])
    ], [(425,280,555,90),(425,340,555,300),(1030,295,1055,295),(425,410,555,540)], (1400,650))
    diagram('state-flow.png', [
        ((30,20,1370,140),'Load required',['Mutations blocked until a successful snapshot load']),
        ((30,200,670,340),'Load success',['Ready for an operation']),
        ((760,200,1370,340),'Load failure',['Error; permission remains invalid']),
        ((30,410,670,550),'Save pending',['Await repository; busy state']),
        ((760,410,1370,550),'Save failure',['Keep accepted records; recover']),
        ((30,620,1370,750),'Save success',['Publish committed collection; update UI'])
    ], [(350,140,350,195),(1060,140,1060,195),(350,340,350,405),(670,480,755,480),(350,550,350,615)], (1400,770))


def blocks(path):
    lines = path.read_text(encoding='utf-8').splitlines(); result=[]; i=0
    while i < len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('#'):
            m=re.match(r'^(#+) (.*)$',line); result.append(('heading',len(m[1]),m[2])); i+=1
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line); result.append(('image',m[1],m[2])); i+=1
        elif line.startswith('Table ') and next((x for x in lines[i+1:] if x.strip()),'').startswith('|'):
            caption=line; i+=1
            while i<len(lines) and not lines[i].strip(): i+=1
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^[-: ]+$', x) for x in cells): rows.append(cells)
                i+=1
            result.append(('table',caption,rows))
        else:
            para=[line]; i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','![','Table ','|')):
                para.append(lines[i].strip()); i+=1
            result.append(('text',' '.join(para).replace('`','')))
    return result


def configure_section(section, numbering=None):
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.top_margin=Cm(3.5); section.bottom_margin=Cm(3)
    section.left_margin=Cm(3.5); section.right_margin=Cm(2)
    section.header_distance=Cm(1.8)
    section.header.is_linked_to_previous=False; section.footer.is_linked_to_previous=False
    for area in (section.header,section.footer):
        for p in area.paragraphs: p.clear()
    p=section.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0); field(p,'PAGE','1')
    if numbering:
        el=OxmlElement('w:pgNumType'); el.set(qn('w:fmt'),numbering); el.set(qn('w:start'),'1')
        section._sectPr.append(el)


def build_docx(all_chapters):
    doc=Document(REF)
    for name in ('Title','Heading 1','Heading 2','Heading 3','Caption'):
        if name not in doc.styles: doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    # Retain source cover paragraphs and embedded logo relationships.
    pars=list(doc.paragraphs)
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr') and child not in [p._p for p in pars[:67]]:
            doc._element.body.remove(child)
    slots={1:'VIETNAM GENERAL CONFEDERATION OF LABOUR',2:'TON DUC THANG UNIVERSITY',3:'FACULTY OF INFORMATION TECHNOLOGY',
        11:'NGUYỄN BÁ HÙNG\nNGUYỄN BẢO LONG',16:TITLE.upper()+'\n'+SUBTITLE.upper(),19:'MIDTERM REPORT',21:'SOFTWARE ENGINEERING',32:'HO CHI MINH CITY 2026',
        34:'VIETNAM GENERAL CONFEDERATION OF LABOUR',35:'TON DUC THANG UNIVERSITY',36:'FACULTY OF INFORMATION TECHNOLOGY',
        45:'NGUYỄN BÁ HÙNG — 523K0006\nNGUYỄN BẢO LONG — 523K0014',50:TITLE.upper()+'\n'+SUBTITLE.upper(),54:'MIDTERM REPORT',56:'SOFTWARE ENGINEERING',59:'Instructor',60:'Mai Văn Mạnh',66:'HO CHI MINH CITY 2026'}
    for i,p in enumerate(pars[:67]):
        if i in (6,40):
            p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(10)
            continue
        p.clear(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.page_break_before=(i==33)
        if i in slots:
            r=p.add_run(slots[i]); r.font.name='Times New Roman'; r.font.size=Pt(14)
            r.font.color.rgb=RGBColor(0,0,0); r.bold=i not in (1,34,59,60,32,66)
            if i in (16,50): p.style='Title'; r.font.size=Pt(22); r.bold=True
            if i in (19,54,21,56): r.font.size=Pt(18)
            p.paragraph_format.line_spacing=1.15
        else:
            p.add_run().font.size=Pt(5); p.paragraph_format.line_spacing=Pt(5)
    # Use two editable spacing slots for verified course metadata.
    pars[24].text='Cross-Platform Mobile App Development — 503107\nClass 23K50201\nSemester 1, Academic Year 2026–2027'
    pars[62].text='Class 23K50201\nCross-Platform Mobile App Development — 503107\nSemester 1, Academic Year 2026–2027'
    for idx in (24,62):
        for r in pars[idx].runs: r.font.name='Times New Roman'; r.font.size=Pt(13)
        pars[idx].paragraph_format.line_spacing=1.25
    for idx in (10,15,18,23,30,44,49,53,58,65): pars[idx].paragraph_format.line_spacing=Pt(15)
    for section in doc.sections:
        for area in (section.header,section.footer,section.first_page_header,section.first_page_footer):
            for p in area.paragraphs: p.clear()
        section.different_first_page_header_footer=False
    styles=doc.styles
    for name in ('Report Body','Report Reference','Report Front','Report Figure','Report Table'):
        if name not in styles: styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    for name in ('Report Body','Report Reference','Report Front','Report Figure','Report Table','Heading 1','Heading 2','Heading 3','Caption','Title'):
        st=styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(13); st.paragraph_format.space_after=Pt(6)
    st=styles['Report Body']; st.paragraph_format.line_spacing=1.5
    st.paragraph_format.first_line_indent=Cm(1.27); st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    st.paragraph_format.space_after=Pt(0)
    st.paragraph_format.widow_control=True
    st=styles['Report Reference']; st.paragraph_format.line_spacing=1.15
    st.paragraph_format.left_indent=Cm(.75); st.paragraph_format.first_line_indent=Cm(-.75)
    st.paragraph_format.space_after=Pt(12)
    for n,size in [('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
        st=styles[n]; st.font.size=Pt(size); st.font.bold=True; st.font.italic=False
        st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.first_line_indent=Pt(0); st.paragraph_format.space_before=Pt(12)
        st.paragraph_format.space_after=Pt(8); st.paragraph_format.keep_with_next=True
        st.paragraph_format.line_spacing=1.15
    styles['Heading 1'].paragraph_format.page_break_before=True
    styles['Report Front'].font.size=Pt(16); styles['Report Front'].font.bold=True
    styles['Report Front'].paragraph_format.space_after=Pt(16)
    styles['Report Front'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    styles['Caption'].font.size=Pt(11); styles['Caption'].font.italic=True
    styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    styles['Caption'].paragraph_format.first_line_indent=Pt(0)
    styles['Caption'].paragraph_format.line_spacing=1
    styles['Report Table'].font.size=Pt(11); styles['Report Table'].paragraph_format.line_spacing=1.1
    styles['Report Table'].paragraph_format.space_after=Pt(5)
    for name in ('TOC 1','TOC 2'):
        if name not in styles: styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        styles[name].font.name='Times New Roman'; styles[name].font.size=Pt(12)
        styles[name].paragraph_format.line_spacing=1.15
        styles[name].paragraph_format.space_after=Pt(2)
    configure_section(doc.add_section(WD_SECTION_START.NEW_PAGE),'lowerRoman')
    doc.add_paragraph('ABSTRACT','Report Front'); doc.add_paragraph(ABSTRACT,'Report Body')
    doc.add_paragraph('Keywords: Flutter; UI automation; integration testing; regression; reproducibility.','Report Body')
    doc.add_page_break(); doc.add_paragraph('CONTENTS','Report Front'); field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')
    doc.add_page_break(); doc.add_paragraph('LIST OF FIGURES','Report Front'); field(doc.add_paragraph(),'TOC \\t "Report Figure,1" \\h \\z')
    doc.add_paragraph('LIST OF TABLES','Report Front'); field(doc.add_paragraph(),'TOC \\t "Report Table Caption,1" \\h \\z')
    if 'Report Table Caption' not in styles: styles.add_style('Report Table Caption',WD_STYLE_TYPE.PARAGRAPH)
    styles['Report Table Caption'].base_style=styles['Caption']
    styles['Report Figure'].base_style=styles['Caption']; styles['Report Figure'].font.size=Pt(11)
    styles['Report Figure'].font.italic=True
    styles['Report Figure'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    styles['Report Table Caption'].paragraph_format.keep_with_next=True
    configure_section(doc.add_section(WD_SECTION_START.NEW_PAGE),'decimal')
    chapter=0; sub=0
    for path in all_chapters+[SOURCE/'references.md',SOURCE/'appendices.md']:
        reference=path.name=='references.md'; appendix=path.name=='appendices.md'
        for b in blocks(path):
            if b[0]=='heading':
                level,title=b[1:]
                if level==1:
                    if not reference and not appendix: chapter+=1; sub=0; title=f'{chapter} {title}'
                elif level==2 and not reference and not appendix: sub+=1; title=f'{chapter}.{sub} {title}'
                p=doc.add_paragraph(title,f'Heading {level}')
                if chapter==1 and level==1: p.paragraph_format.page_break_before=False
            elif b[0]=='text': doc.add_paragraph(b[1],'Report Reference' if reference else 'Report Body')
            elif b[0]=='image':
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                r=p.add_run(); r.add_picture(str(LATEX/b[2]),width=Cm(15.2))
                pic=r._r.xpath('.//wp:docPr')[0]; pic.set('descr',b[1])
                doc.add_paragraph(b[1],'Report Figure')
            elif b[0]=='table':
                doc.add_paragraph(b[1],'Report Table Caption')
                rows=b[2]; table=doc.add_table(rows=1,cols=len(rows[0])); table.autofit=False
                widths=[3.3,5.4,6.8]
                for col,w in zip(table.columns,widths): col.width=Cm(w)
                for j,row in enumerate(rows):
                    cells=table.rows[0].cells if j==0 else table.add_row().cells
                    for k,text in enumerate(row):
                        cells[k].width=Cm(widths[k]); p=cells[k].paragraphs[0]; p.style='Report Table'
                        r=p.add_run(text.upper() if j==0 else text); r.bold=j==0
                        tcPr=cells[k]._tc.get_or_add_tcPr()
                        borders=OxmlElement('w:tcBorders')
                        bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'4'); bottom.set(qn('w:color'),'CBD5E1'); borders.append(bottom); tcPr.append(borders)
                        if j==0:
                            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F1F5F9'); tcPr.append(shade)
                    trPr=cells[0]._tc.getparent().get_or_add_trPr()
                    trPr.append(OxmlElement('w:cantSplit'))
                    if j==0: trPr.append(OxmlElement('w:tblHeader'))
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.core_properties.title=TITLE+' '+SUBTITLE
    doc.core_properties.author='Nguyễn Bá Hùng; Nguyễn Bảo Long'
    doc.core_properties.subject='TaskFlow QA Lab midterm report, course 503107'
    doc.core_properties.comments=''
    doc.save(OUT/'taskflow-report.docx')
    with zipfile.ZipFile(OUT/'taskflow-report.docx') as z, zipfile.ZipFile(REF) as original:
        assert z.read('word/media/image1.png')==original.read('word/media/image1.png')


def tex_escape(s):
    chars={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    # Allow breaking long source paths and URLs without printing extra glyphs.
    parts=re.split(r'(https?://\S+|(?:docs|scripts|integration_test|test|C:)/\S+)',s)
    return ''.join(r'\nolinkurl{'+p+'}' if re.match(r'^(https?://|(?:docs|scripts|integration_test|test|C:)/)',p) else ''.join(chars.get(c,c) for c in p) for p in parts)


def build_tex(all_chapters):
    pre=r'''\documentclass[12pt,a4paper,oneside]{report}
\usepackage{fontspec}
\setmainfont{Times New Roman}
\setsansfont{Arial}
\usepackage[top=3.5cm,bottom=3cm,left=3.5cm,right=2cm,headheight=16pt,headsep=1cm]{geometry}
\usepackage{graphicx,array,longtable,booktabs,caption,titlesec,fancyhdr,xurl,hyperref}
\hypersetup{hidelinks,pdftitle={UI Automation Testing in Flutter},pdfauthor={Nguyễn Bá Hùng and Nguyễn Bảo Long}}
\urlstyle{same}
\makeatletter
\renewcommand\normalsize{\@setfontsize\normalsize{13pt}{19.5pt}}
\makeatother
\AtBeginDocument{\normalsize}
\setlength{\parindent}{1.27cm}
\setlength{\parskip}{0pt}
\setlength{\emergencystretch}{2em}
\widowpenalty=10000\clubpenalty=10000
\pagestyle{fancy}\fancyhf{}\fancyhead[C]{\thepage}\renewcommand{\headrulewidth}{0pt}
\fancypagestyle{plain}{\fancyhf{}\fancyhead[C]{\thepage}\renewcommand{\headrulewidth}{0pt}}
\titleformat{\chapter}[hang]{\bfseries\fontsize{16}{20}\selectfont}{\thechapter}{0.7em}{}
\titlespacing*{\chapter}{0pt}{0pt}{16pt}
\titleformat{\section}{\bfseries\fontsize{14}{18}\selectfont}{\thesection}{0.7em}{}
\titlespacing*{\section}{0pt}{12pt}{6pt}
\setcounter{tocdepth}{1}
\captionsetup{font={small,it},labelfont=it,justification=centering}
\begin{document}\hypersetup{pageanchor=false}
'''
    pages=[]
    for inner in (False,True):
        pages.append(r'\begin{titlepage}\thispagestyle{empty}\centering'+'\n'+
            r'{\fontsize{14}{18}\selectfont VIETNAM GENERAL CONFEDERATION OF LABOUR\\\textbf{TON DUC THANG UNIVERSITY}\\\textbf{FACULTY OF INFORMATION TECHNOLOGY}\par}'+'\n'+
            r'\vspace{0.7cm}\includegraphics[width=3.1cm]{assets/tdtu-logo.png}\par\vspace{0.8cm}'+'\n'+
            r'{\bfseries\fontsize{14}{19}\selectfont NGUYỄN BÁ HÙNG'+(' — 523K0006' if inner else '')+r'\\NGUYỄN BẢO LONG'+(' — 523K0014' if inner else '')+r'\par}\vspace{1cm}'+'\n'+
            r'{\bfseries\fontsize{22}{27}\selectfont UI AUTOMATION TESTING\\IN FLUTTER\\IMPLEMENTING END-TO-END TESTING\par}\vspace{0.8cm}'+'\n'+
            r'{\bfseries\fontsize{18}{23}\selectfont MIDTERM REPORT\\SOFTWARE ENGINEERING\par}\vspace{0.6cm}'+'\n'+
            (r'Instructor: Mai Văn Mạnh\par\vspace{0.35cm}' if inner else '')+'\n'+
            r'Cross-Platform Mobile App Development — 503107\\Class 23K50201\\Semester 1, Academic Year 2026–2027\par\vfill HO CHI MINH CITY 2026\end{titlepage}'+'\n')
    body=[pre]+pages+[r'\pagenumbering{roman}\hypersetup{pageanchor=true}\chapter*{ABSTRACT}',tex_escape(ABSTRACT)+'\n\n',r'\noindent\textit{Keywords: Flutter; UI automation; integration testing; regression; reproducibility.}',r'\tableofcontents\clearpage\listoffigures\listoftables\clearpage\pagenumbering{arabic}']
    append=False
    for path in all_chapters+[SOURCE/'references.md',SOURCE/'appendices.md']:
        refs=path.name=='references.md'
        if path.name=='appendices.md': body.append(r'\appendix'); append=True
        for b in blocks(path):
            if b[0]=='heading':
                level,title=b[1:]
                if append and level==1: title=re.sub(r'^Appendix [A-Z] ','',title)
                if refs: body.append(r'\chapter*{References}\addcontentsline{toc}{chapter}{References}')
                else: body.append('\\'+('chapter' if level==1 else 'section')+'{'+tex_escape(title)+'}')
            elif b[0]=='text':
                text=tex_escape(b[1])
                body.append((r'\noindent\hangindent=0.75cm\hangafter=1 '+text+r'\par\medskip') if refs else text+'\n')
            elif b[0]=='image':
                caption=re.sub(r'^Figure \d+\.\d+ ','',b[1])
                body.append(r'\begin{figure}[htbp]\centering\includegraphics[width=\linewidth]{'+b[2]+r'}\caption{'+tex_escape(caption)+r'}\end{figure}')
            elif b[0]=='table':
                caption=re.sub(r'^Table [\dA-Z]+\.\d+ ','',b[1]); rows=b[2]
                head=' & '.join(r'\textbf{'+tex_escape(c.upper())+'}' for c in rows[0])+r' \\ \midrule'
                body.extend([r'{\fontsize{11}{14}\selectfont\setlength{\tabcolsep}{4pt}\renewcommand{\arraystretch}{1.15}',r'\begin{longtable}{@{}>{\raggedright\arraybackslash}p{0.21\linewidth}>{\raggedright\arraybackslash}p{0.34\linewidth}>{\raggedright\arraybackslash}p{\dimexpr0.45\linewidth-16pt\relax}@{}}',r'\caption{'+tex_escape(caption)+r'}\\\toprule',head,r'\endfirsthead\toprule',head,r'\endhead'])
                for row in rows[1:]: body.append(' & '.join(tex_escape(c) for c in row)+r' \\ \addlinespace[5pt]')
                body.extend([r'\bottomrule\end{longtable}}'])
    body.append(r'\end{document}')
    (LATEX/'taskflow-report.tex').write_text('\n'.join(body),encoding='utf-8')


if __name__=='__main__':
    assert hashlib.sha256(REF.read_bytes()).hexdigest()==REF_HASH, 'Reference template has changed'
    chapters=sorted((SOURCE/'chapters').glob('*.md'))
    assets(); build_docx(chapters); build_tex(chapters)
    print(json.dumps({'docx':str(OUT/'taskflow-report.docx'),'latex':str(LATEX/'taskflow-report.tex'),'main_words':sum(len(p.read_text(encoding='utf-8').split()) for p in chapters)},indent=2))
